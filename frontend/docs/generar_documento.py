"""Genera Documentacion-Asgard-TouristChain.docx y .pdf desde docs/*.md + diagramas."""
import os
import re

BASE = os.path.dirname(os.path.abspath(__file__))
DOCX_OUT = os.path.join(BASE, "Documentacion-Asgard-TouristChain.docx")
PDF_OUT = os.path.join(BASE, "Documentacion-Asgard-TouristChain.pdf")

SECTIONS = [
    "01-empresa.md",
    "02-producto.md",
    "03-manual-usuario.md",
    "04-manual-admin-socio.md",
    "05-api.md",
    "06-instalacion.md",
    "07-datos-prueba.md",
    "08-seguridad-limitaciones.md",
]

NAVY = "1A2A6C"
GOLD = "FFD60A"

# ---------- mini parser markdown ----------
def parse_md(path):
    blocks = []
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    i, buf, in_code = 0, [], False
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(("code", "\n".join(buf))); buf = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            buf.append(line); i += 1
            continue
        if line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif re.match(r"^!\[.*\]\(.*\)$", line.strip()):
            m = re.match(r"^!\[(.*)\]\((.*)\)$", line.strip())
            img = m.group(2)
            if not os.path.isabs(img):
                img = os.path.join(BASE, img)
            blocks.append(("img", (m.group(1), img)))
        elif line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", (header, rows)))
            continue
        elif re.match(r"^(\-|\*) ", line.strip()):
            items = []
            while i < len(lines) and re.match(r"^(\-|\*) ", lines[i].strip()):
                items.append(re.sub(r"^(\-|\*) ", "", lines[i].strip()))
                i += 1
            blocks.append(("ul", items))
            continue
        elif re.match(r"^\d+\. ", line.strip()):
            items = []
            while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                items.append(re.sub(r"^\d+\. ", "", lines[i].strip()))
                i += 1
            blocks.append(("ol", items))
            continue
        elif line.strip() in ("---", "***"):
            blocks.append(("hr", None))
        elif line.strip() == "":
            pass
        else:
            blocks.append(("p", line.strip()))
        i += 1
    return blocks

def split_bold(text):
    """Divide en (texto, negrita)."""
    parts, last = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > last:
            parts.append((text[last:m.start()], False))
        parts.append((m.group(1), True))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], False))
    return parts

def clean(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).replace("`", "")

# ---------- DOCX ----------
def build_docx(all_blocks):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for s in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[s]
        st.font.color.rgb = RGBColor(0x1A, 0x2A, 0x6C)
        st.font.bold = True

    def para(text, bold=False, size=None, color=None, align=None):
        p = doc.add_paragraph()
        for chunk, b in split_bold(text):
            r = p.add_run(chunk)
            r.bold = bold or b
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        if align:
            p.alignment = align
        return p

    # PORTADA
    for _ in range(3):
        doc.add_paragraph()
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = logo.add_run("[  PEGAR AQUÍ EL LOGO DE ASGARD  ]")
    r.bold, r.font.size = True, Pt(14)
    r.font.color.rgb = RGBColor(0x1A, 0x2A, 0x6C)
    para("ASGARD", bold=True, size=40, color=RGBColor(0x1A, 0x2A, 0x6C), align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Plataforma TouristChain — Documentación oficial del sistema",
         size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    para("Fundadores: Cristian, David", align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Directorio y programadores: Lucía, Johan, Iván", align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Versión 1.0 · 2026 · Uso interno y de cliente", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # ÍNDICE
    doc.add_heading("Contenido", level=1)
    for idx, name in enumerate(SECTIONS, 1):
        title = open(os.path.join(BASE, name), encoding="utf-8").readline().lstrip("# ").strip()
        para(f"{idx}. {title}")
    doc.add_page_break()

    for name in SECTIONS:
        for kind, data in parse_md(os.path.join(BASE, name)):
            if kind in ("h1", "h2", "h3"):
                doc.add_heading(clean(data), level=int(kind[1]))
            elif kind == "p":
                para(data)
            elif kind == "code":
                p = doc.add_paragraph()
                r = p.add_run(data)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            elif kind in ("ul", "ol"):
                for item in data:
                    p = doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number")
                    for chunk, b in split_bold(item):
                        rr = p.add_run(chunk)
                        rr.bold = b
            elif kind == "table":
                header, rows = data
                t = doc.add_table(rows=1 + len(rows), cols=len(header))
                t.style = "Light Grid Accent 1"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for j, h in enumerate(header):
                    c = t.rows[0].cells[j]
                    c.text = ""
                    rr = c.paragraphs[0].add_run(clean(h))
                    rr.bold = True
                for irow, row in enumerate(rows, 1):
                    for j in range(len(header)):
                        c = t.rows[irow].cells[j]
                        c.text = ""
                        for chunk, b in split_bold(row[j] if j < len(row) else ""):
                            rr = c.paragraphs[0].add_run(chunk)
                            rr.bold = b
                doc.add_paragraph()
            elif kind == "img":
                alt, img = data
                if os.path.exists(img):
                    doc.add_picture(img, width=Cm(15))
                    last = doc.paragraphs[-1]
                    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para(alt, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            elif kind == "hr":
                doc.add_paragraph("—" * 40)

    # pie en cada sección
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = "Asgard · Plataforma TouristChain · Confidencial"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(DOCX_OUT)
    print("DOCX:", DOCX_OUT)

# ---------- PDF ----------
def build_pdf(all_blocks):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.colors import HexColor, white, black
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image, PageBreak, HRFlowable, ListFlowable,
                                    ListItem)
    from reportlab.lib.units import cm

    navy, gold = HexColor("#1A2A6C"), HexColor("#7a6410")
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=17, textColor=navy, spaceAfter=6)
    st_h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=14, textColor=navy, spaceBefore=10, spaceAfter=4)
    st_h3 = ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=12, textColor=navy, spaceBefore=8, spaceAfter=3)
    st_p = ParagraphStyle("p", fontName="Helvetica", fontSize=10.5, leading=15, alignment=TA_JUSTIFY, spaceAfter=4)
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=9, leading=12)
    st_cellb = ParagraphStyle("cb", fontName="Helvetica-Bold", fontSize=9, leading=12)
    st_code = ParagraphStyle("code", fontName="Courier", fontSize=8.5, leading=12, backColor=HexColor("#F2F4F7"))
    st_cap = ParagraphStyle("cap", fontName="Helvetica-Oblique", fontSize=9, alignment=TA_CENTER, spaceAfter=6)

    def fmt(text):
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
        return t.replace("`", "")

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFillColor(navy)
        canvas.rect(0, 0, LETTER[0], 1.1 * cm, fill=1, stroke=0)
        canvas.setFillColor(white)
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(LETTER[0] / 2, 0.55 * cm, f"Asgard · TouristChain · pág. {doc_.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(PDF_OUT, pagesize=LETTER,
                            topMargin=1.5 * cm, bottomMargin=1.8 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            title="Asgard TouristChain - Documentación",
                            author="Asgard")
    story = []
    # PORTADA
    story += [Spacer(1, 4 * cm),
              Paragraph("[ PEGAR AQUÍ EL LOGO DE ASGARD ]", ParagraphStyle("l", parent=st_p, alignment=TA_CENTER, textColor=navy)),
              Spacer(1, 0.5 * cm),
              Paragraph("ASGARD", ParagraphStyle("t0", fontName="Helvetica-Bold", fontSize=38, textColor=navy, alignment=TA_CENTER)),
              Paragraph("Plataforma TouristChain — Documentación oficial", ParagraphStyle("t1", fontName="Helvetica", fontSize=14, alignment=TA_CENTER)),
              Spacer(1, 1 * cm),
              Paragraph("Fundadores: Cristian, David<br/>Directorio y programadores: Lucía, Johan, Iván",
                        ParagraphStyle("t2", fontName="Helvetica", fontSize=11, alignment=TA_CENTER)),
              Paragraph("Versión 1.0 · 2026 · Uso interno y de cliente", ParagraphStyle("t3", fontName="Helvetica", fontSize=10, alignment=TA_CENTER)),
              PageBreak(),
              Paragraph("Contenido", st_title)]
    for idx, name in enumerate(SECTIONS, 1):
        title = open(os.path.join(BASE, name), encoding="utf-8").readline().lstrip("# ").strip()
        story.append(Paragraph(f"{idx}. {title}", st_p))
    story.append(PageBreak())

    for name in SECTIONS:
        for kind, data in parse_md(os.path.join(BASE, name)):
            if kind == "h1":
                story.append(Paragraph(fmt(data), st_title))
            elif kind == "h2":
                story.append(Paragraph(fmt(data), st_h2))
            elif kind == "h3":
                story.append(Paragraph(fmt(data), st_h3))
            elif kind == "p":
                story.append(Paragraph(fmt(data), st_p))
            elif kind == "code":
                story.append(Paragraph(data.replace("\n", "<br/>"), st_code))
                story.append(Spacer(1, 4))
            elif kind in ("ul", "ol"):
                items = [ListItem(Paragraph(fmt(x), st_p), leftIndent=18) for x in data]
                story.append(ListFlowable(items, bulletType="bullet" if kind == "ul" else "1"))
                story.append(Spacer(1, 4))
            elif kind == "table":
                header, rows = data
                tdata = [[Paragraph(f"<b>{fmt(h)}</b>", st_cellb) for h in header]]
                for row in rows:
                    tdata.append([Paragraph(fmt(row[j] if j < len(row) else ""), st_cell) for j in range(len(header))])
                t = Table(tdata, repeatRows=1)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), navy),
                    ("TEXTCOLOR", (0, 0), (-1, 0), white),
                    ("GRID", (0, 0), (-1, -1), 0.5, navy),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [white, HexColor("#F2F4F7")]),
                ]))
                story.append(t)
                story.append(Spacer(1, 6))
            elif kind == "img":
                alt, img = data
                if os.path.exists(img):
                    story.append(Image(img, width=16 * cm, height=10 * cm, kind="proportional"))
                    story.append(Paragraph(alt, st_cap))
            elif kind == "hr":
                story.append(HRFlowable(width="100%", color=navy))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    print("PDF:", PDF_OUT)

if __name__ == "__main__":
    blocks = {n: parse_md(os.path.join(BASE, n)) for n in SECTIONS}
    build_docx(blocks)
    build_pdf(blocks)
